import io
import json
import os
import random
import time
import uuid
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from flask import Flask, jsonify, request, send_file, send_from_directory
from flask_cors import CORS

app = Flask(__name__, static_folder="static", static_url_path="")
CORS(app)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
DATA_FILE = os.path.join(DATA_DIR, "gastos.json")
ASSINATURAS_FILE = os.path.join(DATA_DIR, "assinaturas.json")
FEATURES_FILE = os.path.join(DATA_DIR, "features.json")

LOCK_TIMEOUT = 5.0
MAX_RETRIES = 15
RETRY_DELAY_MIN = 0.05
RETRY_DELAY_MAX = 0.35

MESES = [
    "Janeiro",
    "Fevereiro",
    "Março",
    "Abril",
    "Maio",
    "Junho",
    "Julho",
    "Agosto",
    "Setembro",
    "Outubro",
    "Novembro",
    "Dezembro",
]

DEFAULT_SECOES = ["Despesas fixas", "Moradia", "Cartões", "Bancos", "Outros"]
DEFAULT_SECOES_RECEITA = ["Receitas", "Salários", "Outras receitas"]

os.makedirs(DATA_DIR, exist_ok=True)


def safe_read_json(filepath):
    if not os.path.exists(filepath):
        return None
    lock_path = filepath + ".lock"
    deadline = time.time() + LOCK_TIMEOUT
    while time.time() < deadline and os.path.exists(lock_path):
        time.sleep(random.uniform(0.02, 0.08))
    for _ in range(MAX_RETRIES):
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_RDWR)
            os.close(fd)
            try:
                with open(filepath, "r", encoding="utf-8-sig") as f:
                    return json.load(f)
            finally:
                if os.path.exists(lock_path):
                    os.remove(lock_path)
        except FileExistsError:
            time.sleep(random.uniform(RETRY_DELAY_MIN, RETRY_DELAY_MAX))
    raise RuntimeError(f"Nao foi possivel ler {filepath}")


def safe_write_json(filepath, data):
    lock_path = filepath + ".lock"
    deadline = time.time() + LOCK_TIMEOUT
    while time.time() < deadline and os.path.exists(lock_path):
        time.sleep(random.uniform(0.02, 0.08))
    for _ in range(MAX_RETRIES):
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_RDWR)
            os.close(fd)
            try:
                with open(filepath, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
                return True
            finally:
                if os.path.exists(lock_path):
                    os.remove(lock_path)
        except FileExistsError:
            time.sleep(random.uniform(RETRY_DELAY_MIN, RETRY_DELAY_MAX))
    raise RuntimeError(f"Nao foi possivel gravar {filepath}")


def default_data():
    return {
        "secoes_despesa": list(DEFAULT_SECOES),
        "secoes_receita": list(DEFAULT_SECOES_RECEITA),
        "anos": [datetime.now().year],
        "lancamentos": [],
        "lixeira": [],
        "meses_revisados": [],
    }


def _log_entry(acao, antes=None, depois=None):
    entry = {"ts": datetime.now().isoformat(timespec="seconds"), "acao": acao}
    if antes is not None:
        entry["antes"] = antes
    if depois is not None:
        entry["depois"] = depois
    return entry


def _snapshot(lanc):
    """Campos relevantes para o log de histórico."""
    return {
        k: lanc.get(k)
        for k in ("descricao", "valor", "secao", "observacao", "tags", "pago", "investido", "tipo", "mes", "ano")
    }


def _ultima_alteracao_ts(lanc):
    historico = lanc.get("historico") or []
    if historico:
        return historico[-1].get("ts") or lanc.get("criado_em")
    return lanc.get("criado_em")


def _legacy_categoria_nome(data, lanc):
    nome = (lanc.get("categoria") or "").strip()
    if nome:
        return nome
    categoria_id = lanc.get("categoria_id")
    if not categoria_id:
        return ""
    for key in ("categorias_receita", "categorias_despesa"):
        for cat in data.get(key, []):
            if cat.get("id") == categoria_id:
                return cat.get("nome") or ""
    return ""


def migrate_lancamentos(data):
    changed = False
    for lanc in data.get("lancamentos", []):
        descricao = (lanc.get("descricao") or "").strip()
        if descricao:
            continue
        legacy = _legacy_categoria_nome(data, lanc)
        if legacy:
            lanc["descricao"] = legacy
            changed = True
    return changed


def migrate_secoes_receita(data):
    """Garante que toda secao usada em receitas exista em secoes_receita."""
    secoes = data.setdefault("secoes_receita", list(DEFAULT_SECOES_RECEITA))
    if not secoes:
        secoes.extend(DEFAULT_SECOES_RECEITA)
    existentes = {s.lower() for s in secoes}
    changed = False
    for lanc in data.get("lancamentos", []):
        if lanc.get("tipo") != "receita":
            continue
        nome = (lanc.get("secao") or "").strip()
        if not nome:
            continue
        if nome.lower() not in existentes:
            secoes.append(nome)
            existentes.add(nome.lower())
            changed = True
    return changed


def get_data():
    data = safe_read_json(DATA_FILE)
    if data is None:
        data = default_data()
        safe_write_json(DATA_FILE, data)
    data.setdefault("secoes_despesa", list(DEFAULT_SECOES))
    data.setdefault("secoes_receita", list(DEFAULT_SECOES_RECEITA))
    data.setdefault("anos", [])
    data.setdefault("lancamentos", [])
    data.setdefault("lixeira", [])
    data.setdefault("meses_revisados", [])
    changed = False
    if migrate_lancamentos(data):
        changed = True
    if migrate_secoes_receita(data):
        changed = True
    if migrate_anos(data):
        changed = True
    if changed:
        safe_write_json(DATA_FILE, data)
    return data


def migrate_anos(data):
    """Garante que todo ano usado em lancamentos exista em data['anos']."""
    anos = data.setdefault("anos", [])
    existentes = set(anos)
    changed = False
    for lanc in data.get("lancamentos", []):
        ano = lanc.get("ano")
        if isinstance(ano, int) and ano not in existentes:
            anos.append(ano)
            existentes.add(ano)
            changed = True
    if not anos:
        anos.append(datetime.now().year)
        changed = True
    return changed


def save_data(data):
    safe_write_json(DATA_FILE, data)


def resolve_descricao(data, lanc):
    descricao = (lanc.get("descricao") or "").strip()
    if descricao:
        return descricao
    legacy = _legacy_categoria_nome(data, lanc)
    return legacy or "Sem descrição"


def normalize_tags(raw):
    if not raw:
        return []
    if isinstance(raw, str):
        raw = [t.strip() for t in raw.split(",")]
    seen = set()
    tags = []
    for tag in raw:
        name = (tag or "").strip()
        if not name:
            continue
        key = name.lower()
        if key in seen:
            continue
        seen.add(key)
        tags.append(name)
    return tags


def collect_all_tags(lancamentos):
    seen = set()
    tags = []
    for lanc in lancamentos:
        for tag in normalize_tags(lanc.get("tags")):
            key = tag.lower()
            if key not in seen:
                seen.add(key)
                tags.append(tag)
    return sorted(tags, key=str.lower)


def filter_lancamentos(lancamentos, ano=None, mes=None, tipo=None, tag=None):
    result = lancamentos
    if ano is not None:
        result = [l for l in result if l.get("ano") == ano]
    if mes is not None:
        result = [l for l in result if l.get("mes") == mes]
    if tipo:
        result = [l for l in result if l.get("tipo") == tipo]
    if tag:
        tag_key = tag.strip().lower()
        result = [
            l
            for l in result
            if any(t.lower() == tag_key for t in normalize_tags(l.get("tags")))
        ]
    return result


def calc_totais(data, ano, mes=None, tag=None):
    lancamentos = filter_lancamentos(data["lancamentos"], ano=ano, mes=mes, tag=tag)
    receitas = [l for l in lancamentos if l.get("tipo") == "receita"]
    despesas = [l for l in lancamentos if l.get("tipo") == "despesa"]
    entrada = sum(l["valor"] for l in receitas)
    entrada_investida = sum(l["valor"] for l in receitas if l.get("investido"))
    saida = sum(l["valor"] for l in despesas)
    saida_paga = sum(l["valor"] for l in despesas if l.get("pago"))
    saida_pendente = saida - saida_paga
    return {
        "entrada": round(entrada, 2),
        "entrada_investida": round(entrada_investida, 2),
        "saida": round(saida, 2),
        "saida_paga": round(saida_paga, 2),
        "saida_pendente": round(saida_pendente, 2),
        "caixa": round(entrada - entrada_investida - saida_paga, 2),
        "liquido": round(entrada - saida, 2),
    }


def _build_sections(grouped, ordem_preferida):
    sections = []
    usadas = set()
    for secao in ordem_preferida:
        items = grouped.get(secao, [])
        if items:
            sections.append(
                {
                    "secao": secao,
                    "itens": items,
                    "total": round(sum(i["valor"] for i in items), 2),
                }
            )
            usadas.add(secao)
    for secao, items in grouped.items():
        if secao in usadas:
            continue
        sections.append(
            {
                "secao": secao,
                "itens": items,
                "total": round(sum(i["valor"] for i in items), 2),
            }
        )
    return sections


def group_by_section(data, lancamentos):
    receitas_por_secao = defaultdict(list)
    despesas_por_secao = defaultdict(list)
    for lanc in lancamentos:
        item = {
            "id": lanc["id"],
            "descricao": resolve_descricao(data, lanc),
            "valor": lanc.get("valor", 0),
            "observacao": lanc.get("observacao", ""),
            "secao": lanc.get("secao") or "Geral",
            "tags": normalize_tags(lanc.get("tags")),
            "pago": bool(lanc.get("pago", False)),
            "investido": bool(lanc.get("investido", False)),
            "criado_em": lanc.get("criado_em"),
            "ultima_alteracao": _ultima_alteracao_ts(lanc),
        }
        if lanc.get("tipo") == "receita":
            receitas_por_secao[item["secao"]].append(item)
        else:
            despesas_por_secao[item["secao"]].append(item)

    receitas_sections = _build_sections(
        receitas_por_secao, data.get("secoes_receita", [])
    )
    despesas_sections = _build_sections(
        despesas_por_secao, data.get("secoes_despesa", [])
    )
    return receitas_sections, despesas_sections


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/api/meta")
def meta():
    return jsonify({"meses": MESES})


def _meses_revisados_do_ano(data, ano):
    return sorted(
        r["mes"]
        for r in data.get("meses_revisados", [])
        if r.get("ano") == ano and isinstance(r.get("mes"), int) and 1 <= r["mes"] <= 12
    )


@app.route("/api/revisao")
def get_revisao():
    ano = request.args.get("ano", type=int)
    if not ano:
        return jsonify({"error": "Parametro 'ano' obrigatorio"}), 400
    data = get_data()
    return jsonify({"ano": ano, "revisados": _meses_revisados_do_ano(data, ano)})


@app.route("/api/revisao/marcar", methods=["POST"])
def set_revisao():
    body = request.get_json(silent=True) or {}
    try:
        ano = int(body.get("ano"))
        mes = int(body.get("mes"))
    except (TypeError, ValueError):
        return jsonify({"error": "ano e mes obrigatorios"}), 400
    if mes < 1 or mes > 12:
        return jsonify({"error": "mes deve ser entre 1 e 12"}), 400
    if "revisado" not in body:
        return jsonify({"error": "Campo 'revisado' obrigatorio"}), 400

    revisado = bool(body["revisado"])
    data = get_data()
    lista = data.setdefault("meses_revisados", [])
    lista[:] = [r for r in lista if not (r.get("ano") == ano and r.get("mes") == mes)]
    if revisado:
        lista.append({"ano": ano, "mes": mes})
    save_data(data)
    return jsonify({"ok": True, "ano": ano, "revisados": _meses_revisados_do_ano(data, ano)})


@app.route("/api/anos")
def list_anos():
    data = get_data()
    derivados = {l.get("ano") for l in data["lancamentos"] if l.get("ano")}
    cadastrados = set(data.get("anos", []))
    anos = sorted(derivados | cadastrados, reverse=True)
    if not anos:
        anos = [datetime.now().year]
    return jsonify({"anos": anos})


@app.route("/api/anos", methods=["POST"])
def create_ano():
    body = request.get_json(silent=True) or {}
    raw = body.get("ano")
    try:
        ano = int(raw)
    except (TypeError, ValueError):
        return jsonify({"error": "Ano invalido"}), 400
    if ano < 1900 or ano > 2200:
        return jsonify({"error": "Ano fora do intervalo permitido (1900-2200)"}), 400

    data = get_data()
    anos = data.setdefault("anos", [])
    if ano in anos:
        return jsonify({"error": "Ano ja cadastrado"}), 409
    anos.append(ano)
    save_data(data)
    todos = sorted(
        set(anos) | {l.get("ano") for l in data["lancamentos"] if l.get("ano")},
        reverse=True,
    )
    return jsonify({"ok": True, "ano": ano, "anos": todos}), 201


@app.route("/api/anos/<int:ano>", methods=["DELETE"])
def delete_ano(ano):
    force = request.args.get("force", "").lower() in ("1", "true", "yes")
    data = get_data()
    lancamentos_ano = [l for l in data["lancamentos"] if l.get("ano") == ano]
    anos = data.setdefault("anos", [])

    if ano not in anos and not lancamentos_ano:
        return jsonify({"error": "Ano nao encontrado"}), 404

    if lancamentos_ano and not force:
        return jsonify(
            {
                "error": "Ano possui lancamentos",
                "lancamentos": len(lancamentos_ano),
                "hint": "Envie ?force=true para excluir o ano e todos os lancamentos.",
            }
        ), 409

    removidos = len(lancamentos_ano)
    if lancamentos_ano:
        data["lancamentos"] = [l for l in data["lancamentos"] if l.get("ano") != ano]
    if ano in anos:
        anos.remove(ano)

    save_data(data)
    return jsonify(
        {
            "ok": True,
            "ano": ano,
            "lancamentos_removidos": removidos,
            "anos": sorted(anos, reverse=True),
        }
    )


def _secoes_key(tipo):
    if tipo == "receita":
        return "secoes_receita"
    if tipo == "despesa":
        return "secoes_despesa"
    return None


@app.route("/api/secoes")
def list_secoes():
    data = get_data()
    return jsonify(
        {
            "secoes": data["secoes_despesa"],
            "secoes_despesa": data["secoes_despesa"],
            "secoes_receita": data["secoes_receita"],
        }
    )


@app.route("/api/secoes", methods=["POST"])
def create_secao():
    body = request.get_json(silent=True) or {}
    tipo = (body.get("tipo") or "despesa").strip()
    nome = (body.get("nome") or "").strip()

    key = _secoes_key(tipo)
    if key is None:
        return jsonify({"error": "Tipo deve ser 'receita' ou 'despesa'"}), 400
    if not nome:
        return jsonify({"error": "Nome da secao obrigatorio"}), 400

    data = get_data()
    secoes = data.setdefault(key, [])
    if any(s.lower() == nome.lower() for s in secoes):
        return jsonify({"error": "Secao ja existe"}), 409

    secoes.append(nome)
    save_data(data)
    return jsonify({"ok": True, "tipo": tipo, "nome": nome, "secoes": secoes}), 201


@app.route("/api/tags")
def list_tags():
    data = get_data()
    return jsonify({"tags": collect_all_tags(data["lancamentos"])})


@app.route("/api/lancamentos")
def list_lancamentos():
    ano = request.args.get("ano", type=int)
    mes = request.args.get("mes", type=int)
    tipo = request.args.get("tipo", "").strip() or None
    tag = request.args.get("tag", "").strip() or None
    if not ano:
        return jsonify({"error": "Parametro 'ano' obrigatorio"}), 400
    data = get_data()
    lancamentos = filter_lancamentos(
        data["lancamentos"], ano=ano, mes=mes, tipo=tipo, tag=tag
    )
    lancamentos.sort(key=lambda l: (l.get("mes", 0), l.get("tipo", ""), resolve_descricao(data, l)))
    enriched = []
    for lanc in lancamentos:
        enriched.append(
            {
                **lanc,
                "descricao": resolve_descricao(data, lanc),
                "tags": normalize_tags(lanc.get("tags")),
            }
        )
    return jsonify(enriched)


@app.route("/api/lancamentos", methods=["POST"])
def create_lancamento():
    body = request.get_json(silent=True) or {}
    ano = body.get("ano")
    mes = body.get("mes")
    tipo = (body.get("tipo") or "").strip()
    valor = body.get("valor")
    descricao = (body.get("descricao") or body.get("categoria") or "").strip()
    secao = (body.get("secao") or "Geral").strip()
    observacao = (body.get("observacao") or "").strip()
    tags = normalize_tags(body.get("tags"))

    if not ano or not mes or tipo not in ("receita", "despesa"):
        return jsonify({"error": "Campos obrigatorios: ano, mes, tipo (receita|despesa)"}), 400
    try:
        valor = round(float(valor), 2)
    except (TypeError, ValueError):
        return jsonify({"error": "Valor invalido"}), 400
    if valor <= 0:
        return jsonify({"error": "Valor deve ser maior que zero"}), 400
    if not descricao:
        return jsonify({"error": "Informe a descricao"}), 400

    data = get_data()

    if tipo == "receita":
        secao_final = secao or "Receitas"
    else:
        secao_final = secao or "Geral"

    lanc = {
        "id": str(uuid.uuid4()),
        "ano": int(ano),
        "mes": int(mes),
        "tipo": tipo,
        "descricao": descricao,
        "secao": secao_final,
        "valor": valor,
        "observacao": observacao,
        "tags": tags,
        "criado_em": datetime.now().isoformat(timespec="seconds"),
        "historico": [],
    }
    lanc["historico"].append(_log_entry("criado", depois=_snapshot(lanc)))
    data["lancamentos"].append(lanc)

    anos = data.setdefault("anos", [])
    if int(ano) not in anos:
        anos.append(int(ano))

    secoes_key = _secoes_key(tipo)
    if secoes_key:
        secoes_lista = data.setdefault(secoes_key, [])
        if not any(s.lower() == secao_final.lower() for s in secoes_lista):
            secoes_lista.append(secao_final)

    save_data(data)
    return jsonify(lanc), 201


@app.route("/api/lancamentos/limpar-mes", methods=["DELETE"])
def delete_mes():
    ano = request.args.get("ano", type=int)
    mes = request.args.get("mes", type=int)
    if not ano:
        return jsonify({"error": "Parametro 'ano' obrigatorio"}), 400
    if not mes or mes < 1 or mes > 12:
        return jsonify({"error": "Parametro 'mes' obrigatorio (1-12)"}), 400

    data = get_data()
    to_delete = [l for l in data["lancamentos"] if l.get("ano") == ano and l.get("mes") == mes]

    now = datetime.now().isoformat(timespec="seconds")
    for lanc in to_delete:
        lanc.setdefault("historico", []).append(_log_entry("excluido", antes=_snapshot(lanc)))
        lanc["excluido_em"] = now

    ids_delete = {l["id"] for l in to_delete}
    data["lancamentos"] = [l for l in data["lancamentos"] if l["id"] not in ids_delete]
    data.setdefault("lixeira", []).extend(to_delete)

    if to_delete:
        save_data(data)

    return jsonify({"ok": True, "removidos": len(to_delete)})


@app.route("/api/lancamentos/<lanc_id>", methods=["PUT"])
def update_lancamento(lanc_id):
    body = request.get_json(silent=True) or {}
    data = get_data()
    lanc = next((l for l in data["lancamentos"] if l.get("id") == lanc_id), None)
    if not lanc:
        return jsonify({"error": "Lancamento nao encontrado"}), 404

    antes = _snapshot(lanc)
    campos_alterados_antes = {}
    campos_alterados_depois = {}

    def _track(campo, valor_novo):
        velho = lanc.get(campo)
        if velho != valor_novo:
            campos_alterados_antes[campo] = velho
            campos_alterados_depois[campo] = valor_novo

    if "valor" in body:
        try:
            valor = round(float(body["valor"]), 2)
        except (TypeError, ValueError):
            return jsonify({"error": "Valor invalido"}), 400
        if valor <= 0:
            return jsonify({"error": "Valor deve ser maior que zero"}), 400
        _track("valor", valor)
        lanc["valor"] = valor

    if "observacao" in body:
        v = (body.get("observacao") or "").strip()
        _track("observacao", v)
        lanc["observacao"] = v

    if "tags" in body:
        v = normalize_tags(body.get("tags"))
        _track("tags", v)
        lanc["tags"] = v

    if "descricao" in body or "categoria" in body:
        descricao = (body.get("descricao") or body.get("categoria") or "").strip()
        if not descricao:
            return jsonify({"error": "Descricao obrigatoria"}), 400
        _track("descricao", descricao)
        lanc["descricao"] = descricao

    if "secao" in body and lanc.get("tipo") in ("receita", "despesa"):
        fallback = "Receitas" if lanc.get("tipo") == "receita" else "Geral"
        nova_secao = (body.get("secao") or fallback).strip() or fallback
        _track("secao", nova_secao)
        lanc["secao"] = nova_secao
        secoes_key = _secoes_key(lanc.get("tipo"))
        if secoes_key:
            secoes_lista = data.setdefault(secoes_key, [])
            if not any(s.lower() == nova_secao.lower() for s in secoes_lista):
                secoes_lista.append(nova_secao)

    if "pago" in body and lanc.get("tipo") == "despesa":
        v = bool(body["pago"])
        _track("pago", v)
        lanc["pago"] = v

    if "investido" in body:
        if lanc.get("tipo") != "receita":
            return jsonify({"error": "Investido aplica-se apenas a receitas"}), 400
        v = bool(body["investido"])
        _track("investido", v)
        lanc["investido"] = v

    if campos_alterados_antes:
        keys = list(campos_alterados_antes.keys())
        if keys == ["pago"] and campos_alterados_depois.get("pago"):
            acao = "pago"
        elif keys == ["pago"]:
            acao = "despago"
        elif keys == ["investido"] and campos_alterados_depois.get("investido"):
            acao = "investido"
        elif keys == ["investido"]:
            acao = "desinvestido"
        else:
            acao = "editado"
        lanc.setdefault("historico", []).append(
            _log_entry(acao, antes=campos_alterados_antes, depois=campos_alterados_depois)
        )

    save_data(data)
    return jsonify(lanc)


@app.route("/api/lancamentos/<lanc_id>/historico")
def get_historico(lanc_id):
    data = get_data()
    lanc = next((l for l in data["lancamentos"] if l.get("id") == lanc_id), None)
    if not lanc:
        return jsonify({"error": "Lancamento nao encontrado"}), 404
    return jsonify({
        "id": lanc_id,
        "descricao": resolve_descricao(data, lanc),
        "historico": list(reversed(lanc.get("historico", []))),
    })


@app.route("/api/lancamentos/<lanc_id>", methods=["DELETE"])
def delete_lancamento(lanc_id):
    data = get_data()
    lanc = next((l for l in data["lancamentos"] if l.get("id") == lanc_id), None)
    if not lanc:
        return jsonify({"error": "Lancamento nao encontrado"}), 404

    lanc.setdefault("historico", []).append(
        _log_entry("excluido", antes=_snapshot(lanc))
    )
    lanc["excluido_em"] = datetime.now().isoformat(timespec="seconds")

    data["lancamentos"] = [l for l in data["lancamentos"] if l.get("id") != lanc_id]
    data.setdefault("lixeira", []).append(lanc)
    save_data(data)
    return jsonify({"ok": True})


@app.route("/api/lixeira")
def list_lixeira():
    data = get_data()
    lixeira = sorted(data.get("lixeira", []), key=lambda l: l.get("excluido_em", ""), reverse=True)
    result = []
    for lanc in lixeira:
        result.append({
            "id": lanc["id"],
            "ano": lanc.get("ano"),
            "mes": lanc.get("mes"),
            "mes_nome": MESES[lanc["mes"] - 1] if lanc.get("mes") and 1 <= lanc["mes"] <= 12 else "",
            "tipo": lanc.get("tipo"),
            "descricao": resolve_descricao(data, lanc),
            "valor": lanc.get("valor"),
            "secao": lanc.get("secao"),
            "excluido_em": lanc.get("excluido_em"),
            "historico": lanc.get("historico", []),
        })
    return jsonify({"lixeira": result, "total": len(result)})


@app.route("/api/lixeira/<lanc_id>/restaurar", methods=["POST"])
def restaurar_lancamento(lanc_id):
    data = get_data()
    lanc = next((l for l in data.get("lixeira", []) if l.get("id") == lanc_id), None)
    if not lanc:
        return jsonify({"error": "Lancamento nao encontrado na lixeira"}), 404

    lanc.pop("excluido_em", None)
    lanc.setdefault("historico", []).append(_log_entry("restaurado"))
    data["lixeira"] = [l for l in data["lixeira"] if l.get("id") != lanc_id]
    data["lancamentos"].append(lanc)

    ano = lanc.get("ano")
    if isinstance(ano, int) and ano not in data.setdefault("anos", []):
        data["anos"].append(ano)

    save_data(data)
    return jsonify({"ok": True, "lancamento": lanc})


@app.route("/api/lixeira/<lanc_id>", methods=["DELETE"])
def delete_permanente(lanc_id):
    data = get_data()
    before = len(data.get("lixeira", []))
    data["lixeira"] = [l for l in data.get("lixeira", []) if l.get("id") != lanc_id]
    if len(data["lixeira"]) == before:
        return jsonify({"error": "Lancamento nao encontrado na lixeira"}), 404
    save_data(data)
    return jsonify({"ok": True})


@app.route("/api/lixeira", methods=["DELETE"])
def esvaziar_lixeira():
    data = get_data()
    total = len(data.get("lixeira", []))
    data["lixeira"] = []
    save_data(data)
    return jsonify({"ok": True, "removidos": total})


@app.route("/api/resumo")
def resumo():
    ano = request.args.get("ano", type=int)
    mes = request.args.get("mes", type=int)
    if not ano:
        return jsonify({"error": "Parametro 'ano' obrigatorio"}), 400

    tag = request.args.get("tag", "").strip() or None
    data = get_data()
    if mes:
        lancamentos = filter_lancamentos(data["lancamentos"], ano=ano, mes=mes, tag=tag)
        receitas_secoes, despesas_secoes = group_by_section(data, lancamentos)
        totais = calc_totais(data, ano, mes, tag=tag)
        receitas_flat = [item for s in receitas_secoes for item in s["itens"]]
        return jsonify(
            {
                "ano": ano,
                "mes": mes,
                "mes_nome": MESES[mes - 1] if 1 <= mes <= 12 else "",
                "totais": totais,
                "receitas": receitas_flat,
                "receitas_por_secao": receitas_secoes,
                "despesas_por_secao": despesas_secoes,
            }
        )

    mensal = []
    for m in range(1, 13):
        totais = calc_totais(data, ano, m)
        mensal.append({"mes": m, "mes_nome": MESES[m - 1], **totais})

    totais_ano = calc_totais(data, ano)
    por_descricao = defaultdict(lambda: {"receita": 0.0, "despesa": 0.0})
    for lanc in filter_lancamentos(data["lancamentos"], ano=ano):
        nome = resolve_descricao(data, lanc)
        por_descricao[nome][lanc["tipo"]] += lanc.get("valor", 0)

    ranking_despesas = sorted(
        [
            {"descricao": nome, "total": round(vals["despesa"], 2)}
            for nome, vals in por_descricao.items()
            if vals["despesa"] > 0
        ],
        key=lambda x: x["total"],
        reverse=True,
    )[:10]

    return jsonify(
        {
            "ano": ano,
            "totais_ano": totais_ano,
            "mensal": mensal,
            "ranking_despesas": ranking_despesas,
        }
    )


TEMPLATE_HEADERS = [
    "Ano", "Mes", "Tipo", "Descricao", "Secao", "Valor", "Observacao", "Tags",
]

TEMPLATE_EXEMPLOS = [
    [2026, 1, "receita", "Salário",      "Receitas",       5000.00, "",            "salario,mensal"],
    [2026, 1, "despesa", "Condomínio",   "Despesas fixas",  850.00, "vencimento dia 10", "casa,fixo"],
    [2026, 1, "despesa", "Cartão Nubank","Cartões",        1200.00, "",            "cartao"],
]

MESES_MAP = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "março": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9,
    "outubro": 10, "novembro": 11, "dezembro": 12,
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}


def _parse_mes(raw):
    if raw is None or raw == "":
        return None
    if isinstance(raw, int):
        return raw if 1 <= raw <= 12 else None
    try:
        n = int(str(raw).strip())
        if 1 <= n <= 12:
            return n
    except ValueError:
        pass
    key = str(raw).strip().lower()
    return MESES_MAP.get(key)


def _parse_tipo(raw):
    if not raw:
        return None
    key = str(raw).strip().lower()
    if key in ("receita", "entrada", "r"):
        return "receita"
    if key in ("despesa", "saida", "saída", "d"):
        return "despesa"
    return None


def _parse_valor(raw):
    if raw is None or raw == "":
        return None
    if isinstance(raw, (int, float)):
        return round(float(raw), 2)
    text = str(raw).strip().replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return round(float(text), 2)
    except ValueError:
        return None


@app.route("/api/template-excel")
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lancamentos"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="D49B3B", end_color="D49B3B", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    for col, name in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, exemplo in enumerate(TEMPLATE_EXEMPLOS, 2):
        for col_idx, val in enumerate(exemplo, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    larguras = [8, 6, 12, 28, 22, 12, 30, 22]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Instrucoes")
    instrucoes = [
        ["Coluna",      "Obrigatorio", "Formato / Exemplo"],
        ["Ano",         "Sim",         "Numero inteiro (ex.: 2026)"],
        ["Mes",         "Sim",         "Numero 1-12 ou nome do mes (ex.: Janeiro)"],
        ["Tipo",        "Sim",         "receita ou despesa"],
        ["Descricao",   "Sim",         "Texto livre (ex.: Condominio)"],
        ["Secao",       "Nao",         "Despesas fixas, Cartoes, Receitas, etc."],
        ["Valor",       "Sim",         "Numero positivo (ex.: 850,00)"],
        ["Observacao",  "Nao",         "Texto livre"],
        ["Tags",        "Nao",         "Separadas por virgula (ex.: casa,fixo)"],
    ]
    for r, linha in enumerate(instrucoes, 1):
        for c, val in enumerate(linha, 1):
            cell = ws2.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="modelo-gastos.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/lancamentos/import-excel", methods=["POST"])
def import_excel():
    if "arquivo" not in request.files:
        return jsonify({"error": "Envie o arquivo no campo 'arquivo'"}), 400
    f = request.files["arquivo"]
    if not f.filename:
        return jsonify({"error": "Arquivo vazio"}), 400
    if not f.filename.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "Use um arquivo .xlsx"}), 400

    try:
        wb = openpyxl.load_workbook(f, data_only=True)
    except Exception as exc:
        return jsonify({"error": f"Nao foi possivel abrir o arquivo: {exc}"}), 400

    sheet_name = request.form.get("aba") or "Lancamentos"
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "Planilha vazia"}), 400

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    def col(name):
        candidatos = {
            "ano":        ["ano"],
            "mes":        ["mes", "mês"],
            "tipo":       ["tipo"],
            "descricao":  ["descricao", "descrição", "descricao do lancamento"],
            "secao":      ["secao", "seção", "categoria"],
            "valor":      ["valor", "valor (r$)", "preco"],
            "observacao": ["observacao", "observação", "obs"],
            "tags":       ["tags", "etiquetas"],
        }
        for alvo in candidatos[name]:
            if alvo in header:
                return header.index(alvo)
        return -1

    idx = {k: col(k) for k in
           ["ano", "mes", "tipo", "descricao", "secao", "valor", "observacao", "tags"]}
    obrigatorios = ["ano", "mes", "tipo", "descricao", "valor"]
    faltando = [c for c in obrigatorios if idx[c] < 0]
    if faltando:
        return jsonify({
            "error": f"Cabecalho invalido. Colunas obrigatorias ausentes: {', '.join(faltando)}",
            "esperado": TEMPLATE_HEADERS,
        }), 400

    data = get_data()
    criados = 0
    erros = []

    def get(row, key):
        i = idx[key]
        return row[i] if i >= 0 and i < len(row) else None

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(c in (None, "") for c in row):
            continue

        try:
            ano = int(get(row, "ano")) if get(row, "ano") not in (None, "") else None
        except (TypeError, ValueError):
            ano = None
        mes = _parse_mes(get(row, "mes"))
        tipo = _parse_tipo(get(row, "tipo"))
        descricao = (str(get(row, "descricao")).strip() if get(row, "descricao") else "")
        secao = (str(get(row, "secao")).strip() if get(row, "secao") else "")
        valor = _parse_valor(get(row, "valor"))
        observacao = (str(get(row, "observacao")).strip() if get(row, "observacao") else "")
        tags_raw = get(row, "tags")

        problemas = []
        if not ano or ano < 1900 or ano > 2200:
            problemas.append("Ano invalido")
        if not mes:
            problemas.append("Mes invalido")
        if not tipo:
            problemas.append("Tipo deve ser 'receita' ou 'despesa'")
        if not descricao:
            problemas.append("Descricao obrigatoria")
        if valor is None or valor <= 0:
            problemas.append("Valor invalido")
        if problemas:
            erros.append({"linha": row_num, "problemas": problemas})
            continue

        if not secao:
            secao = "Receitas" if tipo == "receita" else "Geral"

        tags = normalize_tags(
            tags_raw if isinstance(tags_raw, list) else str(tags_raw or "")
        )

        lanc = {
            "id": str(uuid.uuid4()),
            "ano": int(ano),
            "mes": int(mes),
            "tipo": tipo,
            "descricao": descricao,
            "secao": secao,
            "valor": valor,
            "observacao": observacao,
            "tags": tags,
            "criado_em": datetime.now().isoformat(timespec="seconds"),
        }
        data["lancamentos"].append(lanc)

        anos = data.setdefault("anos", [])
        if lanc["ano"] not in anos:
            anos.append(lanc["ano"])

        secoes_key = _secoes_key(tipo)
        if secoes_key:
            secoes_lista = data.setdefault(secoes_key, [])
            if not any(s.lower() == secao.lower() for s in secoes_lista):
                secoes_lista.append(secao)

        criados += 1

    save_data(data)
    return jsonify({
        "ok": True,
        "criados": criados,
        "erros": erros,
        "total_linhas": len(rows) - 1,
    })


# ── Assinaturas / custos recorrentes (cartão) ─────────────────────


def default_assinaturas_data():
    return {"cartoes": [], "assinaturas": []}


def get_assinaturas_data():
    data = safe_read_json(ASSINATURAS_FILE)
    if data is None:
        data = default_assinaturas_data()
        safe_write_json(ASSINATURAS_FILE, data)
    data.setdefault("cartoes", [])
    data.setdefault("assinaturas", [])
    return data


def save_assinaturas_data(data):
    safe_write_json(ASSINATURAS_FILE, data)


def _parse_iso_date(value, field_name, required=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError(f"Campo '{field_name}' obrigatorio")
        return None
    s = str(value).strip()[:10]
    try:
        datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"Data invalida em '{field_name}' (use AAAA-MM-DD)")
    return s


def _assinatura_snapshot(item):
    return {
        k: item.get(k)
        for k in (
            "descricao",
            "data_inicio",
            "data_fim",
            "valor_mensal",
            "cartao",
        )
    }


def _assinatura_ultima_alteracao(item):
    historico = item.get("historico") or []
    if historico:
        return historico[-1].get("ts")
    return item.get("criado_em")


def _assinatura_ativa(item, ref_date=None):
    fim = item.get("data_fim")
    if not fim:
        return True
    ref = ref_date or datetime.now().strftime("%Y-%m-%d")
    return fim >= ref


def _enrich_assinatura(item):
    return {
        **item,
        "ultima_alteracao": _assinatura_ultima_alteracao(item),
        "ativa": _assinatura_ativa(item),
    }


def _registrar_cartao(data, nome):
    cartao = (nome or "").strip()
    if not cartao:
        return
    cartoes = data.setdefault("cartoes", [])
    if not any(c.lower() == cartao.lower() for c in cartoes):
        cartoes.append(cartao)
        cartoes.sort(key=str.lower)


def _validar_datas_assinatura(data_inicio, data_fim):
    inicio = _parse_iso_date(data_inicio, "data_inicio", required=True)
    fim = _parse_iso_date(data_fim, "data_fim", required=False)
    if fim and fim < inicio:
        raise ValueError("data_fim nao pode ser anterior a data_inicio")
    return inicio, fim


@app.route("/api/assinaturas/cartoes")
def list_cartoes_assinaturas():
    data = get_assinaturas_data()
    return jsonify({"cartoes": sorted(data.get("cartoes", []), key=str.lower)})


@app.route("/api/assinaturas")
def list_assinaturas():
    cartao = (request.args.get("cartao") or "").strip() or None
    apenas_ativas = request.args.get("ativas", "").lower() in ("1", "true", "sim")
    data = get_assinaturas_data()
    items = list(data.get("assinaturas", []))
    if cartao:
        items = [a for a in items if (a.get("cartao") or "").lower() == cartao.lower()]
    if apenas_ativas:
        items = [a for a in items if _assinatura_ativa(a)]
    items.sort(key=lambda a: ((a.get("cartao") or "").lower(), (a.get("descricao") or "").lower()))
    enriched = [_enrich_assinatura(a) for a in items]
    total_mensal_ativas = round(
        sum(float(a.get("valor_mensal") or 0) for a in enriched if a.get("ativa")), 2
    )
    return jsonify({"assinaturas": enriched, "total_mensal_ativas": total_mensal_ativas})


@app.route("/api/assinaturas", methods=["POST"])
def create_assinatura():
    body = request.get_json(silent=True) or {}
    descricao = (body.get("descricao") or "").strip()
    cartao = (body.get("cartao") or "").strip()
    if not descricao:
        return jsonify({"error": "Informe a descricao"}), 400
    if not cartao:
        return jsonify({"error": "Informe o cartao de credito"}), 400
    try:
        data_inicio, data_fim = _validar_datas_assinatura(
            body.get("data_inicio"), body.get("data_fim")
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    try:
        valor_mensal = round(float(body.get("valor_mensal")), 2)
    except (TypeError, ValueError):
        return jsonify({"error": "Valor mensal invalido"}), 400
    if valor_mensal <= 0:
        return jsonify({"error": "Valor mensal deve ser maior que zero"}), 400

    data = get_assinaturas_data()
    now = datetime.now().isoformat(timespec="seconds")
    item = {
        "id": str(uuid.uuid4()),
        "descricao": descricao,
        "data_inicio": data_inicio,
        "data_fim": data_fim,
        "valor_mensal": valor_mensal,
        "cartao": cartao,
        "criado_em": now,
        "historico": [],
    }
    item["historico"].append(_log_entry("criado", depois=_assinatura_snapshot(item)))
    data["assinaturas"].append(item)
    _registrar_cartao(data, cartao)
    save_assinaturas_data(data)
    return jsonify(_enrich_assinatura(item)), 201


@app.route("/api/assinaturas/<item_id>", methods=["PUT"])
def update_assinatura(item_id):
    body = request.get_json(silent=True) or {}
    data = get_assinaturas_data()
    item = next((a for a in data["assinaturas"] if a.get("id") == item_id), None)
    if not item:
        return jsonify({"error": "Assinatura nao encontrada"}), 404

    antes = _assinatura_snapshot(item)
    campos_alterados_antes = {}
    campos_alterados_depois = {}

    def _track(campo, valor_novo):
        velho = item.get(campo)
        if velho != valor_novo:
            campos_alterados_antes[campo] = velho
            campos_alterados_depois[campo] = valor_novo

    if "descricao" in body:
        descricao = (body.get("descricao") or "").strip()
        if not descricao:
            return jsonify({"error": "Informe a descricao"}), 400
        _track("descricao", descricao)
        item["descricao"] = descricao

    if "cartao" in body:
        cartao = (body.get("cartao") or "").strip()
        if not cartao:
            return jsonify({"error": "Informe o cartao de credito"}), 400
        _track("cartao", cartao)
        item["cartao"] = cartao
        _registrar_cartao(data, cartao)

    if "valor_mensal" in body:
        try:
            valor_mensal = round(float(body["valor_mensal"]), 2)
        except (TypeError, ValueError):
            return jsonify({"error": "Valor mensal invalido"}), 400
        if valor_mensal <= 0:
            return jsonify({"error": "Valor mensal deve ser maior que zero"}), 400
        _track("valor_mensal", valor_mensal)
        item["valor_mensal"] = valor_mensal

    if "data_inicio" in body or "data_fim" in body:
        try:
            data_inicio, data_fim = _validar_datas_assinatura(
                body.get("data_inicio", item.get("data_inicio")),
                body.get("data_fim", item.get("data_fim")),
            )
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        _track("data_inicio", data_inicio)
        _track("data_fim", data_fim)
        item["data_inicio"] = data_inicio
        item["data_fim"] = data_fim

    if campos_alterados_antes:
        item.setdefault("historico", []).append(
            _log_entry("editado", antes=campos_alterados_antes, depois=campos_alterados_depois)
        )

    save_assinaturas_data(data)
    return jsonify(_enrich_assinatura(item))


@app.route("/api/assinaturas/<item_id>/historico")
def get_assinatura_historico(item_id):
    data = get_assinaturas_data()
    item = next((a for a in data["assinaturas"] if a.get("id") == item_id), None)
    if not item:
        return jsonify({"error": "Assinatura nao encontrada"}), 404
    return jsonify({
        "id": item_id,
        "descricao": item.get("descricao") or "",
        "historico": list(reversed(item.get("historico", []))),
    })


@app.route("/api/assinaturas/<item_id>", methods=["DELETE"])
def delete_assinatura(item_id):
    data = get_assinaturas_data()
    item = next((a for a in data["assinaturas"] if a.get("id") == item_id), None)
    if not item:
        return jsonify({"error": "Assinatura nao encontrada"}), 404
    data["assinaturas"] = [a for a in data["assinaturas"] if a.get("id") != item_id]
    save_assinaturas_data(data)
    return jsonify({"ok": True})


# ── Features implementadas (changelog) ────────────────────────────


def default_features_data():
    return {"features": []}


def get_features_data():
    data = safe_read_json(FEATURES_FILE)
    if data is None:
        data = default_features_data()
        safe_write_json(FEATURES_FILE, data)
    data.setdefault("features", [])
    return data


def _feature_sort_key(item):
    return item.get("implementado_em") or ""


@app.route("/api/features")
def list_features():
    data = get_features_data()
    items = list(data.get("features", []))
    items.sort(key=_feature_sort_key, reverse=True)
    return jsonify({"features": items, "total": len(items)})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5001"))
    print("=" * 60)
    print("  Controle de Gastos Pessoais")
    print(f"  Dados: {DATA_FILE}")
    print(f"  Assinaturas: {ASSINATURAS_FILE}")
    print(f"  Features: {FEATURES_FILE}")
    print(f"  Servidor: http://localhost:{port}")
    print("=" * 60)
    app.run(host="0.0.0.0", port=port, debug=False)
